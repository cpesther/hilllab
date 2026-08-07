# Christopher Esther, Hill Lab, 7/17/2026
from pptx import Presentation

def parse_powerpoint(full_path):

    """Powerpoint file parser"""

    try:
        prs = Presentation(full_path)
        words = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    if shape.text.strip():
                        words.append(shape.text)

        return words

    except Exception:
        return None
